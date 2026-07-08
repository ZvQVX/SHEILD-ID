"""
Generate SHIELD-ID final report as a Word .docx (run: python generate_shield_final_report_docx.py).
Requires: pip install python-docx
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_para(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(11)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def build() -> Document:
    doc = Document()
    title = doc.add_heading("SHIELD-ID — Final Project Report", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_para(
        doc,
        "Continuous behavioral biometrics, Active Directory integration, "
        "ML-assisted anomaly scoring, and encrypted audit logging.",
    )
    doc.add_paragraph()

    # --- Structured summary (list) ---
    add_heading(doc, "Structured summary", 1)
    bullets = [
        "Purpose: Ongoing keystroke timing and mouse movement signals combined with enterprise identity (Active Directory) and encrypted audit storage.",
        "Sensitive data: Behavioral telemetry is treated as highly sensitive (HSPI). Plaintext audit lines are not stored in the primary path; each line is encrypted with Python cryptography Fernet before append.",
        "Data at rest (security_audit.enc): Encapsulation (encrypt before write); C = Fernet.encrypt(key, plaintext); append-only log of one token per line.",
        "Data in transit: HTTPS on port 8443 (Flask run_https); LDAP/LDAPS to AD configurable via environment (e.g. AD_PORT, AD_USE_SSL).",
        "Key management (shield.key): Fernet.generate_key() on first run; offline decryption via decrypt_audit.py; decrypt failures indicate wrong key or corrupted ciphertext.",
        "Machine learning: sklearn RandomForestClassifier (100 trees when no pkl); optional biometric_rf_model.pkl via joblib; features dwell (ms) and velocity (px/s); output anomaly score 0–100 from predict_proba or heuristic fallback.",
        "Dataset / collection: Desktop path uses pynput listeners and encrypts each line to security_audit.enc; web UI posts telemetry and derives dwell from keydown intervals; no in-repo training pipeline—document labeling and training if using a real model.",
        "Identity: AD integration and JWT-backed API access; RBAC depends on directory configuration and claim mapping.",
        "Conclusion: Layered prototype—directory auth, behavioral telemetry, ML scoring, TLS transport, Fernet-encrypted append-only audit with supervisor decryption.",
    ]
    add_bullets(doc, bullets)

    doc.add_page_break()

    # --- Narrative sections ---
    add_heading(doc, "The cryptographic framework", 1)
    add_para(
        doc,
        "In SHIELD-ID, behavioral data such as keystroke timing and mouse activity is treated as "
        "Highly Sensitive Personal Information (HSPI). The system uses a “Vault” architecture built on "
        "symmetric authenticated encryption via Python’s Fernet (cryptography library), which uses "
        "AES-based encryption with an HMAC for integrity.",
    )
    add_heading(doc, "Data at rest: the security_audit.enc repository", 2)
    add_bullets(
        doc,
        [
            "Encapsulation: When the system logs a session or telemetry, the payload is encrypted before being written to disk.",
            "Encryption process: A key K derived from shield.key is used with plaintext payload P to produce token C (Fernet encrypt).",
            "Immutable logging: security_audit.enc is an append-only encrypted log. Each line is a separate Fernet token, so undetected tampering of individual records is difficult without the key.",
        ],
    )
    add_heading(doc, "Data in transit: secure transport layer", 2)
    add_bullets(
        doc,
        [
            "Protocol: Communication between the biometric monitor / browser and the backend API uses HTTPS (TLS).",
            "Port configuration: The backend may use port 8443 for HTTPS, providing an encrypted channel on the local network.",
            "Directory transport: LDAP or LDAPS to the configured Active Directory host should be used per deployment policy; verify TLS settings in production.",
        ],
    )
    add_heading(doc, "Key management and shield.key", 2)
    add_bullets(
        doc,
        [
            "The Vault depends on shield.key, the Fernet master key file for all secured log lines.",
            "Separation of duties: The key file is stored outside any “database” table and should be restricted to administrators.",
            "Decryption gate: To read encrypted logs, an administrator runs decrypt_audit.py in the project folder, which applies Fernet.decrypt per line.",
            "Tamper detection: Invalid ciphertext or wrong key causes decryption to fail (authenticated encryption).",
        ],
    )
    add_para(
        doc,
        "This layered approach keeps data protected at rest and in transit while providing clear safeguards against casual unauthorized changes.",
    )

    add_heading(doc, "Machine learning model", 1)
    add_para(
        doc,
        "The anomaly engine is implemented in ml_engine.py as sklearn.ensemble.RandomForestClassifier "
        "with n_estimators=100 when no persisted model is loaded. If biometric_rf_model.pkl exists, it is "
        "loaded with joblib. The API exposes dwell and velocity; predict_proba drives a 0–100 style score, "
        "with a normalized heuristic fallback if the model is missing or errors.",
    )

    add_heading(doc, "How the dataset was built or collected", 1)
    add_para(
        doc,
        "Desktop collection (main.py): pynput keyboard and mouse listeners record key release dwell time "
        "(ms) and mouse velocity (px/s). Each scored event is formatted as a timestamped line, encrypted "
        "by SecurityVault, and appended to security_audit.enc.",
    )
    add_para(
        doc,
        "Web / dashboard path: The React application sends authenticated telemetry to the backend and "
        "can derive dwell from keydown timing. Demo mode may add simulated variance to some UI metrics; "
        "state clearly in your defense whether results are measured or illustrative.",
    )
    add_para(
        doc,
        "Supervised training: This repository does not ship a training script that calls fit() on labeled "
        "impostor/genuine data. For a production narrative, describe your labeling scheme, sample counts, "
        "train/validation split, and metrics (e.g. FAR/FRR) once available.",
    )

    add_heading(doc, "Figures and tables — labels and captions", 1)
    add_para(doc, "Use numbered figures and tables in order of appearance. Suggested captions:", bold=True)

    table = doc.add_table(rows=1, cols=3)
    hdr = table.rows[0].cells
    hdr[0].text = "ID"
    hdr[1].text = "Type"
    hdr[2].text = "Caption"
    rows = [
        (
            "Figure 1",
            "Architecture / data-flow",
            "End-to-end data flow: biometric capture, TLS to API, Random Forest scoring, Fernet append to security_audit.enc.",
        ),
        (
            "Figure 2",
            "Sequence diagram",
            "Sequence of AD authentication, JWT use, and telemetry or audit posts to the backend.",
        ),
        (
            "Figure 3",
            "Dashboard screenshot",
            "Operator dashboard: dwell time, mouse velocity, anomaly score, and session log stream.",
        ),
        (
            "Figure 4",
            "Alert / lockdown",
            "Alert state when anomaly score exceeds the configured threshold (strict vs relaxed mode).",
        ),
        (
            "Table 1",
            "Crypto components",
            "Cryptographic components: shield.key, security_audit.enc, decrypt_audit.py.",
        ),
        (
            "Table 2",
            "ML schema",
            "Model inputs dwell (ms), velocity (px/s); output anomaly_score (0–100).",
        ),
        (
            "Table 3",
            "Network parameters",
            "HTTPS port 8443; LDAP 389 vs LDAPS 636; SSL verify flags for deployment.",
        ),
    ]
    for rid, typ, cap in rows:
        row = table.add_row().cells
        row[0].text = rid
        row[1].text = typ
        row[2].text = cap

    add_heading(doc, "Conclusion", 1)
    add_para(
        doc,
        "This phase bridges traditional identity management and behavioral biometrics: Active Directory "
        "integration with configurable LDAPS, JWT-protected APIs, Random Forest–based anomaly scoring, "
        "RBAC as enforced by your directory and API policies, and immutable-style encrypted audit logging "
        "via Fernet. Continuous, non-intrusive authentication is viable as a prototype and can be hardened "
        "for production with explicit TLS verification, training data governance, and operational key management.",
    )

    doc.add_page_break()
    add_heading(doc, "Reviewer comments and feedback", 1)
    feedback = [
        "Script name: Use decrypt_audit.py consistently (not audit_decrypter.py) to match the repository.",
        "Terminology: Prefer “AES-based authenticated encryption (Fernet)” over implying “AES-128 only.”",
        "TLS 1.3: Only claim TLS 1.3 if verified against your OpenSSL/server build; otherwise say “TLS (HTTPS).”",
        "Pinning: Lab deployments may use AD_INSECURE_SKIP_VERIFY; do not equate that with certificate pinning unless you implement it.",
        "Dataset rigor: Add labeling, volumes, splits, and FAR/FRR when you have real study data; disclose simulated UI metrics in demo mode.",
        "Immutability: Append-only Fernet lines deter casual edits; a privileged host could still replace the file—mention WORM or remote logging if required.",
        "Privacy: Add retention, lawful basis, and access controls for HSPI if the report is used in compliance contexts.",
        "Figures: Every chart needs an in-figure title and a caption stating scenario and time range.",
    ]
    add_bullets(doc, feedback)

    add_para(doc, "Generated by generate_shield_final_report_docx.py from the shield-pulse-monitor project.")
    return doc


def main() -> None:
    out = Path(__file__).resolve().parent / "SHIELD-ID-Final-Report.docx"
    doc = build()
    doc.save(out)
    print(f"Wrote: {out}")


if __name__ == "__main__":
    main()
